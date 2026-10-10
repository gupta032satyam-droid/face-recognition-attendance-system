"""Reports page: browse the Excel attendance sheet, search, sort, export CSV."""
import csv
import datetime as dt
import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import config
from gui_common import BG, CARD, FONT, MUTED, cell_text, make_card, open_excel


class ReportsPage(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=BG)
        self.app = app
        self.header, self.rows = [], []
        self._sort_rev = {}
        self._build()

    def _build(self):
        card = make_card(self)
        card.pack(fill="both", expand=True)

        top = tk.Frame(card, bg=CARD)
        top.pack(fill="x", padx=14, pady=12)
        tk.Label(top, text="Search", bg=CARD, fg=MUTED, font=(FONT, 10)).pack(side="left")
        self.search = tk.StringVar()
        self.search.trace_add("write", lambda *_: self._render())
        ttk.Entry(top, textvariable=self.search, width=28).pack(side="left", padx=8)
        ttk.Button(top, text="Today",
                   command=lambda: self.search.set(dt.date.today().isoformat())).pack(side="left")
        ttk.Button(top, text="Clear", command=lambda: self.search.set("")).pack(side="left", padx=6)
        self.count = tk.StringVar()
        tk.Label(top, textvariable=self.count, bg=CARD, fg=MUTED).pack(side="left", padx=10)
        ttk.Button(top, text="Open Excel", command=open_excel).pack(side="right")
        ttk.Button(top, text="Export CSV", style="Accent.TButton",
                   command=self.export_csv).pack(side="right", padx=8)
        ttk.Button(top, text="Refresh", command=self.load).pack(side="right")

        wrap = tk.Frame(card, bg=CARD)
        wrap.pack(fill="both", expand=True, padx=14, pady=(0, 12))
        self.table = ttk.Treeview(wrap, show="headings")
        ys = ttk.Scrollbar(wrap, orient="vertical", command=self.table.yview)
        xs = ttk.Scrollbar(wrap, orient="horizontal", command=self.table.xview)
        self.table.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        ys.pack(side="right", fill="y")
        xs.pack(side="bottom", fill="x")
        self.table.pack(fill="both", expand=True)

    def on_show(self):
        self.load()

    def load(self):
        self.header, self.rows = [], []
        if os.path.exists(config.ATTENDANCE_XLSX):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(config.ATTENDANCE_XLSX, read_only=True, data_only=True)
                data = [[cell_text(c) for c in r] for r in wb.active.iter_rows(values_only=True)]
                wb.close()
                data = [r for r in data if any(r)]
                if data:
                    self.header = [h or f"Col {i + 1}" for i, h in enumerate(data[0])]
                    n = len(self.header)
                    self.rows = [(r + [""] * n)[:n] for r in data[1:]]
            except Exception as exc:
                messagebox.showerror("Report", f"Could not read the Excel file:\n{exc}")
        self.table["columns"] = self.header
        for i, h in enumerate(self.header):
            self.table.heading(h, text=h, command=lambda i=i: self._sort(i))
            self.table.column(h, width=150, anchor="w")
        self._render()

    def _filtered(self):
        q = self.search.get().strip().lower()
        if not q:
            return self.rows
        return [r for r in self.rows if q in " ".join(r).lower()]

    def _render(self):
        self.table.delete(*self.table.get_children())
        rows = self._filtered()
        for r in rows:
            self.table.insert("", "end", values=r)
        self.count.set(f"{len(rows)} record(s)")

    def _sort(self, idx):
        rev = self._sort_rev.get(idx, False)
        self.rows.sort(key=lambda r: r[idx].lower(), reverse=rev)
        self._sort_rev[idx] = not rev
        self._render()

    def export_csv(self):
        if not self.header:
            messagebox.showinfo("Export", "Nothing to export yet.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv",
                                            filetypes=[("CSV file", "*.csv")],
                                            initialfile="attendance_export.csv")
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(self.header)
            w.writerows(self._filtered())
        self.app.set_status(f"Exported to {path}")
