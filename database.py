"""Storage layer: students in CSV, attendance in Excel (one sheet per date)."""
import csv
import os
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook

import config

FIELDS = ["label", "roll_no", "name", "email"]
EMAIL_RE = re.compile(r"^[\w.+-]+@[\w-]+(\.[\w-]+)+$")


def init_storage():
    os.makedirs(config.DATASET_DIR, exist_ok=True)
    if not os.path.exists(config.STUDENTS_CSV):
        with open(config.STUDENTS_CSV, "w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=FIELDS).writeheader()


def get_students():
    """Return list of dicts: label (int), roll_no, name, email."""
    init_storage()
    with open(config.STUDENTS_CSV, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["label"] = int(r["label"])
    return rows


def get_student_by_label(label):
    for s in get_students():
        if s["label"] == label:
            return s
    return None


def add_student(roll_no, name, email):
    """Register a new student and return the numeric label used by the model."""
    roll_no, name, email = roll_no.strip(), name.strip(), email.strip()
    if not (roll_no and name and email):
        raise ValueError("Roll number, name and email are all required.")
    if not EMAIL_RE.match(email):
        raise ValueError("Please enter a valid email address.")
    students = get_students()
    if any(s["roll_no"].lower() == roll_no.lower() for s in students):
        raise ValueError(f"Roll number {roll_no} is already enrolled.")
    label = max([s["label"] for s in students], default=0) + 1
    with open(config.STUDENTS_CSV, "a", newline="", encoding="utf-8") as f:
        csv.DictWriter(f, fieldnames=FIELDS).writerow(
            {"label": label, "roll_no": roll_no, "name": name, "email": email})
    return label


def remove_student(label):
    """Used to roll back an enrollment that was cancelled midway."""
    students = [s for s in get_students() if s["label"] != label]
    with open(config.STUDENTS_CSV, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(students)


def _open_workbook():
    if os.path.exists(config.ATTENDANCE_XLSX):
        return load_workbook(config.ATTENDANCE_XLSX)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _today_sheet(wb):
    today = datetime.now().strftime("%Y-%m-%d")
    if today in wb.sheetnames:
        return wb[today]
    ws = wb.create_sheet(today)
    ws.append(["Roll No", "Name", "Email", "Time"])
    for col, width in zip("ABCD", (14, 26, 32, 12)):
        ws.column_dimensions[col].width = width
    return ws


def mark_attendance(student):
    """Mark a student present today. Returns True if newly marked, False if duplicate."""
    init_storage()
    wb = _open_workbook()
    ws = _today_sheet(wb)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == student["roll_no"]:
            return False
    ws.append([student["roll_no"], student["name"], student["email"],
               datetime.now().strftime("%H:%M:%S")])
    wb.save(config.ATTENDANCE_XLSX)   # raises PermissionError if Excel has the file open
    return True


def todays_attendance():
    if not os.path.exists(config.ATTENDANCE_XLSX):
        return []
    wb = load_workbook(config.ATTENDANCE_XLSX, read_only=True)
    today = datetime.now().strftime("%Y-%m-%d")
    if today not in wb.sheetnames:
        return []
    return [r for r in wb[today].iter_rows(min_row=2, values_only=True) if r[0]]
